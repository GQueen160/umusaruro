"""Phase 1 evidence contract for Umusaruro. Regenerates the DOCX and the PDF."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from fpdf import FPDF

DOCX = "/home/spider/Projects/Hackathon-NISR/docs/Umusaruro-evidence-contract.docx"
PDF = "/home/spider/Projects/Hackathon-NISR/docs/Umusaruro-evidence-contract.pdf"

NAVY = RGBColor(0x1C, 0x2A, 0x3A)
INK = RGBColor(0x20, 0x24, 0x2A)
MUTED = RGBColor(0x5C, 0x67, 0x72)


def add_hyperlink(paragraph, text, url):
    r_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1A56C4")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    text_el = OxmlElement("w:t")
    text_el.set(qn("xml:space"), "preserve")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def font(run, size=11, bold=False, color=INK):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = "Calibri"


def heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = NAVY
    return p


def body(doc, text, size=11):
    p = doc.add_paragraph()
    font(p.add_run(text), size=size)
    p.paragraph_format.space_after = Pt(8)
    return p


def link_line(doc, label, url):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.35)
    font(p.add_run(label), size=11, bold=True)
    p.add_run("\n")
    add_hyperlink(p, url, url)
    return p


def shade(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "1C2A3A")
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)
    for p in cell.paragraphs:
        for run in p.runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
            run.font.size = Pt(8)
            run.font.name = "Calibri"


def table(doc, headers, rows):
    grid = doc.add_table(rows=1 + len(rows), cols=len(headers))
    grid.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = grid.rows[0].cells[i]
        cell.text = header
        shade(cell)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = grid.rows[r + 1].cells[c]
            cell.text = value
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8)
                    run.font.name = "Calibri"
    doc.add_paragraph()


def build_docx():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.7)
        section.bottom_margin = Cm(1.7)
        section.left_margin = Cm(1.7)
        section.right_margin = Cm(1.7)

    title = doc.add_paragraph()
    font(title.add_run("Umusaruro"), size=26, bold=True, color=NAVY)
    sub = doc.add_paragraph()
    font(
        sub.add_run("Evidence contract  ·  Phase 1  ·  15 October 2026"),
        size=12,
        color=MUTED,
    )

    body(
        doc,
        "This document locks the product. A number may appear on a screen only when this contract names its survey, table, and year. A blank cell, or a dash, stays blank. It is not zero. The page is for a district agronomist and a planner at the Ministry of Agriculture and Animal Resources. They read it. They do not type figures into it.",
    )
    body(
        doc,
        "The cells below were read on 15 October 2026 from the Season A 2026 PDF and, for Season B, from the NISR publication page and the Season B workbook. Phase 2 still requires a second person to check every cell against the printed page before the file is trusted. This draft was prepared with model assistance. The competition requires that disclosure in the submission.",
    )
    body(
        doc,
        "Repository: https://gitlab.com/mukunzijonathan/umusaruro. The project is private. This file was written in the local folder because that remote asked for a username and could not be cloned from this machine.",
    )

    heading(doc, "1. The product", 1)
    body(
        doc,
        "Umusaruro, the Kinyarwanda word for harvest, is an agricultural advisor for the cultivation journey, arranged by crop. For each crop the Season A 2026 report prints, the crop page shows district yield, and, where that crop is named, national improved seed and national post-harvest loss. The district map still shows maize yield against the PSTA 5 milestones, then the lever the all-crop tables support: irrigation, or improved seed, with inorganic fertilizer named when that table lags. Maize is the milestone crop. The other crops are not scored against that milestone. It is the agriculture track of the NISR 2026 Big Data Hackathon. A submission is a public Git repository, a deployed application, and documentation. The deadline is 23:59 on 30 October 2026.",
    )

    heading(doc, "2. The milestones", 1)
    body(
        doc,
        "PSTA 5 measures maize productivity with the Seasonal Agricultural Survey, in tonnes per hectare. The results-framework header is Baseline (2023), then five annual columns. Season A 2026 maize, all farms, is 1,994 kg/ha, which is 1.99 t/ha after dividing by 1,000 and rounding half up to two decimals. That figure is under the 2025/26 milestone of 2.06 and under the 2026/27 milestone of 2.26. Parliament’s “2.26” is the 2026/27 column. The page names the year.",
    )
    table(
        doc,
        ["Column", "Maize, t/ha", "Source"],
        [
            ["Baseline (2023)", "1.70", "PSTA 5 results framework"],
            ["2024/25", "1.87", "PSTA 5 results framework"],
            ["2025/26", "2.06", "PSTA 5 results framework"],
            ["2026/27", "2.26", "PSTA 5 results framework"],
            ["2027/28", "2.49", "PSTA 5 results framework"],
            ["2028/29", "2.74", "PSTA 5 results framework"],
        ],
    )

    heading(doc, "3. What Season A 2026 may show", 1)
    body(
        doc,
        "The base file is the Season A 2026 report. It is the latest report that prints district maize yield and district irrigation by farm type together, and Season A is the main maize season. It is not the whole year. The screen stores the figure as printed. Kilograms per hectare are shown as tonnes only by dividing by 1,000 and rounding half up to two decimals.",
    )
    table(
        doc,
        ["What", "Table", "Level"],
        [
            ["Yield of each crop Table 13 prints, kg/ha", "13", "District and national. Maize also has SSF and LSF rows."],
            ["Large-farm yield of each crop, kg/ha", "14", "District and national. A dash stays blank."],
            ["Improved seed, all crops, % of farmers", "20", "District, by farm type"],
            ["Source of improved seed", "22", "District"],
            ["Inorganic fertilizer, % of farmers", "25", "District, by farm type"],
            ["Irrigation, % of farmers", "31", "District, by farm type"],
            ["Improved seed by crop, % of seed", "21", "National only. Not copied onto a district."],
            ["Post-harvest loss by crop, % of use", "16", "National only. No district store."],
        ],
    )
    body(
        doc,
        "National maize, Table 13: smallholders 1,897 kg/ha (1.90 t/ha), all farms 1,994 kg/ha (1.99 t/ha), large farms 4,337 kg/ha (4.34 t/ha). Irrigation, Table 31, percentage of farmers: overall 14.4, smallholders 12.6, large farms 60.3. Improved seed, Table 20, all crops, percentage of farmers: overall 37.0, smallholders 35.2, large farms 82.6. Table 21 is a different statistic: maize seed is 64.3% improved and 35.7% traditional, with no district column and no farm-size column. Bush bean improved seed is 0.5%. Beans in Table 13 are 700 kg/ha nationally and 542 kg/ha on large farms. The page does not tell a district to close a bean gap toward large farms.",
    )
    body(
        doc,
        "Table 20 is not the maize seed rate. The 35.2 against 82.6 comparison is all crops. A card that names seed says “all crops”.",
    )

    heading(doc, "4. Districts checked on 15 October 2026", 1)
    body(
        doc,
        "In Table 13, only Nyagatare (2,887 kg/ha, 2.89 t/ha) and Kirehe (2,515 kg/ha, 2.52 t/ha) are at or above 2.26 t/ha. The next three are Rulindo 2,096 kg/ha (2.10), Muhanga 2,081 (2.08), and Burera 2,042 (2.04). Nyagatare and Kirehe are not the reason the national figure is 1.99.",
    )
    table(
        doc,
        ["District", "Maize, Table 13", "Against 2.26", "SSF irrigating, Table 31", "LSF irrigating"],
        [
            ["Nyagatare", "2,887 kg/ha (2.89 t/ha)", "Above", "6.4%", "50.0%"],
            ["Kirehe", "2,515 kg/ha (2.52 t/ha)", "Above", "14.9%", "78.6%"],
            ["Kamonyi", "1,863 kg/ha (1.86 t/ha)", "Below", "22.4%", "93.3%"],
        ],
    )
    body(
        doc,
        "Kamonyi’s large-farm maize yield is 6,470 kg/ha (Table 14), which is 6.47 t/ha. Nyagatare’s large-farm maize yield is 5,292 kg/ha.",
    )
    body(
        doc,
        "All-crop improved seed, Table 20, smallholders against large farms: Nyagatare 53.8% against 92.9% (gap 39.1 points); Kamonyi 45.3% against 100% (gap 54.7); Kirehe 56.0% against 92.9% (gap 36.9); Rwamagana 53.6% against 78.1% (gap 24.5); Ngoma 44.9% against 57.1% (gap 12.2); Kicukiro 38.7% against 60.0% (gap 21.3); Gakenke 26.4% against 55.6% (gap 29.2).",
    )
    body(
        doc,
        "Inorganic fertilizer, Table 25, percentage of farmers: Kamonyi smallholders 47.2% against large farms 93.3% (gap 46.1); Nyagatare 67.9% against 91.7% (gap 23.8); Kirehe 76.0% against 92.9% (gap 16.9); Gakenke 85.1% against 66.7%, so smallholders are ahead and fertilizer is not named.",
    )

    heading(doc, "5. Rules the application may apply", 1)
    body(
        doc,
        "These constants are the team’s. They are not NISR thresholds. They live in one module, and the method page prints them. Changing a constant changes who appears in the data file.",
    )
    body(
        doc,
        "Irrigation is a candidate only when the large-farm rate minus the smallholder rate is at least 30 points and the smallholder rate is under 15%. It becomes the action only when the district’s CHIRPS total for September 2025 through February 2026 is also below the median of the 30 districts. Seed is named when the Table 20 farmer gap is at least 25 points. If irrigation and seed both fire, the larger gap is the action, unless seed leads by less than 5 points, in which case irrigation stays. Fertilizer is named on the same card when the Table 25 smallholder rate lags the large-farm rate by at least 15 points. A blank large-farm cell cannot fire a gap.",
    )
    body(
        doc,
        "This contract does not lock a rainfall total. Phase 2 computes CHIRPS and writes the district total, the number of grid cells, and the source file. Until that file exists, the ten districts below are candidates, not the action list. A claim that rainfall already drops the list to seven is not in this contract.",
    )
    table(
        doc,
        ["Candidate", "SSF %", "LSF %", "Gap, points"],
        [
            ["Kicukiro", "9.9", "80.0", "70.1"],
            ["Kayonza", "12.8", "78.7", "65.9"],
            ["Kirehe", "14.9", "78.6", "63.7"],
            ["Gatsibo", "8.1", "63.6", "55.5"],
            ["Ngoma", "4.7", "50.0", "45.3"],
            ["Nyagatare", "6.4", "50.0", "43.6"],
            ["Rwamagana", "10.4", "53.7", "43.3"],
            ["Karongi", "10.8", "50.0", "39.2"],
            ["Nyamasheke", "8.8", "46.2", "37.4"],
            ["Nyamagabe", "8.8", "42.9", "34.1"],
        ],
    )
    body(
        doc,
        "Worked cards, before rainfall. Nyagatare: irrigation gap 43.6 and seed gap 39.1 both clear their lines, so irrigation is the action; fertilizer lag 23.8 also clears 15 and is named. The card does not say the yield is failing 2.26. Kamonyi: smallholder irrigation is 22.4%, so irrigation does not fire; the action is all-crop seed (gap 54.7), and fertilizer (gap 46.1) is named. Kirehe: irrigation gap 63.7 beats seed gap 36.9, so irrigation is the action if rainfall passes; fertilizer lag 16.9 is named. Gakenke: irrigation gap 28.7 does not fire; seed gap 29.2 does; fertilizer is not named. Ngoma: seed gap 12.2 does not fire, so the candidate action is irrigation, and fertilizer lag 18.2 is named.",
    )
    body(
        doc,
        "Boundaries that sit next to the constants on the method page. Kirehe smallholder irrigation is 14.9% against the 15% line. Muhanga is 15.2% and does not enter the candidate list. Gakenke’s irrigation gap is 28.7 points against 30. Rwamagana’s seed gap is 24.5 points against 25. Kirehe’s fertilizer lag is 16.9 points against 15, so it is named. One tenth of a point changes each of these.",
    )

    heading(doc, "6. What this contract refuses", 1)
    body(
        doc,
        "No district store. Table 16 puts maize post-harvest loss at 0.3% of use and other usage at 0.2%, one national row. A loss figure from another year is a different series and is not pasted onto this table.",
    )
    body(
        doc,
        "No maize variety. The 2025 plot questionnaire records improved, traditional, or both. It does not record a variety name. No satisfaction score. The closest item is a production-status reason, including inappropriate seeds, lack of improved seeds, unfertile soil, and good quality of seeds. That reason is not in the published district tables. It may be used only from microdata, after the file reproduces a published national rate, and it is labeled as a reason.",
    )
    body(
        doc,
        "No learned score of seed against soil and fertilizer. The Seasonal Agricultural Survey does not carry a soil test on the plot. The Rwanda Soil Information System, launched by the Rwanda Agriculture and Animal Resources Development Board in November 2024, recommends fertilizer from soil. Maize recommendations were still in testing at that launch. Umusaruro does not replace that system and does not train a model on it.",
    )
    body(
        doc,
        "Season B does not overwrite a Season A cell. Season C is not averaged into Season A. The irrigated-hectare path in NST2 is not re-estimated. A national rate is not copied onto a district. Rainfall does not replace a NISR cell.",
    )

    heading(doc, "7. Season B 2026, as a check", 1)
    body(
        doc,
        "NISR released Season B 2026 on 22 September 2026. Collection ran from 19 April to 28 June 2026, in 1,200 segments and 505 large-scale farmers. The publication page states: improved seed 20.9% of farmers, 19.4% of smallholders, 65% of large farms; organic fertilizer 81.3%, 81.7%, and 71.3%; inorganic fertilizer 57.5%, 56.7%, and 80.2%; pesticides 39.1%, 37.7%, and 80%; irrigation 13.1% of farmers. The workbook’s national improved-seed row (Table 18) is 20.9, 19.36, and 64.98, which rounds to the page. Its national irrigation row (Table 29) is 13.08, 11.47, and 56.23. Maize yield in workbook Table 6 is 1.29 t/ha, from the cell 1.292, with production 122,370 tonnes. That is not the Season A yield of 1.99.",
    )
    body(
        doc,
        "Season B table numbers are not the Season A table numbers. Seed is Table 18, fertilizer is Table 23, irrigation is Table 29. The footer of workbook Table 29 still reads “SAS 2025”. No Season B irrigation cell goes on a screen until a second person resolves that footer. The page’s 13.1% sentence is the figure this contract relies on.",
    )

    heading(doc, "8. Links", 1)
    links = [
        ("NISR 2026 Big Data Hackathon", "https://statistics.gov.rw/about/hackathon/2026-hackathon-competition"),
        ("SAS Season A 2026, full report", "https://statistics.gov.rw/sites/default/files/documents/2026-04/SAS%202026%20Season%20A%20_Final%20Report.pdf"),
        ("SAS Season B 2026, publication page", "https://statistics.gov.rw/statistical-publications/agriculture-environment/seasonal-agricultural-survey-season-b-2026"),
        ("SAS Season B 2026, main report", "https://statistics.gov.rw/sites/default/files/documents/2026-09/SAS%202026B_Main%20report.pdf"),
        ("SAS Season B 2026, English summary", "https://statistics.gov.rw/sites/default/files/documents/2026-09/SAS%202026B_Summary%20in%20English.pdf"),
        ("SAS Season B 2026, workbook", "https://statistics.gov.rw/sites/default/files/documents/2026-09/Tables_2026%20Season%20B.xlsx"),
        ("SAS 2025 annual report", "https://statistics.gov.rw/sites/default/files/documents/2025-12/SAS%202025%20Final%20report.pdf"),
        ("SAS 2025 microdata catalogue", "https://microdata.statistics.gov.rw/index.php/catalog/124"),
        ("SAS 2025 plot questionnaire", "https://www.microdata.statistics.gov.rw/index.php/catalog/124/download/1114"),
        ("PSTA 5, abridged results framework", "https://www.minagri.gov.rw/fileadmin/user_upload/Minagri/Publications/Policies_and_strategies/PSTA_5_Abridged_version_Final.pdf"),
        ("NST2 2024–2029, abridged", "https://www.minecofin.gov.rw/fileadmin/user_upload/Minecofin/Publications/STRATEGIES/NST_2/NST2_2024-2029_Abridged.pdf"),
        ("Resolutions of the 20th Umushyikirano, 5–6 February 2026", "https://www.umushyikirano.gov.rw/fileadmin/user_upload/Umushyikirano/Imyanzuro/Resolutions_20th_UMUSHYIKIRANO_-_17_Feb_26.pdf"),
        ("Agricultural Household Survey 2024, main report", "https://statistics.gov.rw/sites/default/files/documents/2025-10/AHS%202024_Main%20report.pdf"),
        ("Agricultural Household Survey 2024, survey page", "https://statistics.gov.rw/data-sources/surveys/Agricultural-Household-Survey/agricultural-household-survey-2024"),
        ("CHIRPS rainfall", "https://www.chc.ucsb.edu/data/chirps"),
        ("RAB note on the Rwanda Soil Information System, 29 November 2024", "https://www.rab.gov.rw/1-1/news-details/rab-launches-rwanda-digital-soil-information-system-to-improve-agricultural-productivity"),
        ("SAS Season A 2024, earlier report, not the base file", "https://statistics.gov.rw/sites/default/files/documents/2024-10/SAS_2024%20Season%20A%20Final%20Report%20%282%29_compressed.pdf"),
        ("GitLab repository", "https://gitlab.com/mukunzijonathan/umusaruro"),
    ]
    for label, url in links:
        link_line(doc, label, url)

    body(
        doc,
        "Secondary. These describe a public meeting. They are not a survey cell.",
        size=11,
    )
    link_line(
        doc,
        "KT Press, 23 March 2026, Gako meeting on agricultural follow-through",
        "https://www.ktpress.rw/2026/03/kagame-criticizes-slow-progress-in-agriculture-urges-shift-from-promises-to-action/",
    )
    link_line(
        doc,
        "AllAfrica, 29 June 2026, Parliament and the maize yield. A news report, not the Hansard.",
        "https://allafrica.com/stories/202606300076.html",
    )

    heading(doc, "9. Phases", 1)
    table(
        doc,
        ["Phase", "Done when"],
        [
            ["1. This contract", "Both teammates accept the links, the checked cells, and the refusals. No screen exists yet."],
            ["2. Data file", "Tables 13, 14, 20, 22, 25, and 31 are transcribed. A second person checks every cell. Blanks stay blank. The file reproduces 1,897, 1,994, 4,337, 12.6, and 60.3. CHIRPS totals are computed and stored with the grid-cell count."],
            ["3. Rules", "One module holds 30, 15, 25, 5, and 15. Tests cover Nyagatare, Kirehe, Kamonyi, Gakenke, Ngoma, and a blank large-farm cell (Gicumbi irrigation, or Karongi maize in Table 14)."],
            ["4. Application", "Django. A map, a district page, and a method page. A statistic cannot render without survey, table, and year."],
            ["5. Season B check", "Season B loads beside Season A, with its own table numbers, and cannot overwrite Season A. The Table 29 footer is resolved first."],
            ["6. Submission", "Public repository, deployed link, method page, and a disclosure of model assistance. Deadline 30 October 2026, 23:59."],
        ],
    )
    body(
        doc,
        "Phase 2 starts when both teammates accept this contract. Microdata is a later pass, and only if a weighted file reproduces a published national rate such as 64.3% or 35.2%. Until then the PDF tables are the numbers.",
    )
    doc.save(DOCX)


class Doc(FPDF):
    def footer(self):
        self.set_y(-12)
        self.set_font("Lib", "", 8)
        self.set_text_color(90, 98, 108)
        self.cell(0, 6, "Umusaruro  ·  Evidence contract  ·  Phase 1  ·  15 October 2026", align="L")


def pdf_h(pdf, text):
    pdf.ln(1)
    pdf.set_font("Lib", "B", 12)
    pdf.set_text_color(28, 42, 58)
    pdf.multi_cell(0, 5.5, text)
    pdf.ln(1)


def pdf_p(pdf, text):
    pdf.set_font("Lib", "", 9)
    pdf.set_text_color(32, 36, 42)
    pdf.multi_cell(0, 4.2, text, align="L")
    pdf.ln(1.2)


def pdf_header_row(pdf, headers, widths):
    pdf.set_font("Lib", "B", 7.5)
    pdf.set_fill_color(28, 42, 58)
    pdf.set_text_color(255, 255, 255)
    for header, width in zip(headers, widths):
        pdf.cell(width, 6, header, border=0, fill=True)
    pdf.ln(6)


def pdf_table(pdf, headers, rows, widths):
    line_h = 3.6
    pdf.set_auto_page_break(auto=False)
    pdf_header_row(pdf, headers, widths)
    for i, row in enumerate(rows):
        pdf.set_font("Lib", "", 7.5)
        pdf.set_text_color(32, 36, 42)
        heights = []
        for text, width in zip(row, widths):
            lines = pdf.multi_cell(width - 2, line_h, text, dry_run=True, output="LINES", align="L")
            heights.append(max(1, len(lines)) * line_h + 1.5)
        height = max(heights)
        if pdf.get_y() + height > 278:
            pdf.add_page()
            pdf_header_row(pdf, headers, widths)
        x0, y0 = pdf.get_x(), pdf.get_y()
        x = x0
        for text, width in zip(row, widths):
            pdf.set_fill_color(244, 246, 248) if i % 2 else pdf.set_fill_color(255, 255, 255)
            pdf.rect(x, y0, width, height, style="F")
            pdf.set_xy(x + 1, y0 + 0.6)
            pdf.multi_cell(width - 2, line_h, text, align="L")
            x += width
        pdf.set_xy(x0, y0 + height)
    pdf.ln(2)
    pdf.set_auto_page_break(auto=True, margin=14)


def build_pdf():
    pdf = Doc(format="A4", unit="mm")
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.add_font("Lib", "", "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf")
    pdf.add_font("Lib", "B", "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf")
    pdf.add_page()
    pdf.set_left_margin(14)
    pdf.set_right_margin(14)
    pdf.set_font("Lib", "B", 18)
    pdf.set_text_color(28, 42, 58)
    pdf.cell(0, 8, "Umusaruro", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Lib", "", 9)
    pdf.set_text_color(92, 103, 114)
    pdf.cell(0, 5, "Evidence contract   ·   Phase 1   ·   15 October 2026", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf_p(
        pdf,
        "A number appears on a screen only when this contract names its survey, table, and year. A blank or a dash stays blank. It is not zero. Readers are a district agronomist and a MINAGRI planner. Cells were read on 15 October 2026 from the Season A 2026 PDF. A second person still checks every cell in Phase 2. Model assistance was used to draft this file. Repository: https://gitlab.com/mukunzijonathan/umusaruro",
    )
    pdf_h(pdf, "1. The product")
    pdf_p(
        pdf,
        "Umusaruro is an agricultural advisor for the cultivation journey, arranged by crop. Crop pages show Season A 2026 district yield for every crop Table 13 prints, plus national seed (Table 21) and national post-harvest loss (Table 16) where those tables name the crop. The district map shows maize yield against the PSTA 5 milestones, then irrigation or all-crop improved seed, and names inorganic fertilizer when that table lags. Other crops are not scored against the maize milestone. Competition deadline: 23:59 on 30 October 2026. The submission is a public Git repository, a deployed application, and documentation.",
    )
    pdf_h(pdf, "2. PSTA 5 maize milestones, tonnes per hectare")
    pdf_table(
        pdf,
        ["Baseline (2023)", "2024/25", "2025/26", "2026/27", "2027/28", "2028/29"],
        [["1.70", "1.87", "2.06", "2.26", "2.49", "2.74"]],
        [32, 28, 28, 28, 28, 38],
    )
    pdf_p(
        pdf,
        "Season A 2026, all farms, is 1,994 kg/ha, shown as 1.99 t/ha. That is under 2.06 and under 2.26. The figure 2.26 is the 2026/27 column. Source: PSTA 5 abridged results framework. Display rule: divide kg/ha by 1,000 and round half up to two decimals.",
    )
    pdf_h(pdf, "3. Cells Season A 2026 may show")
    pdf_table(
        pdf,
        ["What", "Table", "Level"],
        [
            ["Each crop's yield, kg/ha", "13", "District and national"],
            ["Each crop's large-farm yield", "14", "District and national. A dash stays blank."],
            ["Improved seed, all crops", "20", "District, by farm type"],
            ["Source of improved seed", "22", "District"],
            ["Inorganic fertilizer", "25", "District, by farm type"],
            ["Irrigation, % of farmers", "31", "District, by farm type"],
            ["Improved seed by crop", "21", "National. Not copied onto a district."],
            ["Post-harvest loss by crop", "16", "National. No district store."],
        ],
        [62, 18, 102],
    )
    pdf_p(
        pdf,
        "National maize: smallholders 1,897 kg/ha (1.90 t/ha), all farms 1,994 (1.99), large farms 4,337 (4.34). Irrigation: 14.4% overall, 12.6% of smallholders, 60.3% of large farms. All-crop improved seed, percentage of farmers: 37.0, 35.2, and 82.6. Table 21 maize seed, 64.3% improved, has no district split. Bush bean improved seed is 0.5%. Beans: 700 kg/ha nationally, 542 kg/ha on large farms. Table 20 is not the maize seed rate.",
    )
    pdf_h(pdf, "4. Three districts, and the irrigation candidates")
    pdf_table(
        pdf,
        ["District", "Maize kg/ha", "t/ha", "Vs 2.26", "SSF irrig.", "LSF irrig."],
        [
            ["Nyagatare", "2,887", "2.89", "Above", "6.4%", "50.0%"],
            ["Kirehe", "2,515", "2.52", "Above", "14.9%", "78.6%"],
            ["Kamonyi", "1,863", "1.86", "Below", "22.4%", "93.3%"],
        ],
        [32, 28, 22, 24, 28, 48],
    )
    pdf_p(
        pdf,
        "Only those two districts are at or above 2.26. Next are Rulindo 2,096 kg/ha, Muhanga 2,081, and Burera 2,042. Kamonyi large-farm maize is 6,470 kg/ha. Seed gaps, Table 20: Nyagatare 39.1 points, Kamonyi 54.7, Kirehe 36.9, Gakenke 29.2, Rwamagana 24.5, Ngoma 12.2. Fertilizer gaps, Table 25: Kamonyi 46.1, Nyagatare 23.8, Kirehe 16.9. In Gakenke, smallholders lead large farms, 85.1% against 66.7%.",
    )
    pdf_p(
        pdf,
        "Team rules, not NISR thresholds. Irrigation candidate: gap at least 30 points and smallholder irrigation under 15%, then a CHIRPS gate for September 2025–February 2026 below the 30-district median. Seed: Table 20 gap at least 25 points. If both fire, the larger gap wins unless seed leads by under 5 points. Fertilizer is named when Table 25 lags by at least 15 points. Rainfall totals are not locked here. The ten rows below are candidates, not the final action list.",
    )
    if pdf.get_y() > 200:
        pdf.add_page()
    pdf_table(
        pdf,
        ["Candidate", "SSF %", "LSF %", "Gap"],
        [
            ["Kicukiro", "9.9", "80.0", "70.1"],
            ["Kayonza", "12.8", "78.7", "65.9"],
            ["Kirehe", "14.9", "78.6", "63.7"],
            ["Gatsibo", "8.1", "63.6", "55.5"],
            ["Ngoma", "4.7", "50.0", "45.3"],
            ["Nyagatare", "6.4", "50.0", "43.6"],
            ["Rwamagana", "10.4", "53.7", "43.3"],
            ["Karongi", "10.8", "50.0", "39.2"],
            ["Nyamasheke", "8.8", "46.2", "37.4"],
            ["Nyamagabe", "8.8", "42.9", "34.1"],
        ],
        [50, 40, 40, 52],
    )
    pdf_p(
        pdf,
        "Before rainfall: Nyagatare’s action is irrigation, and fertilizer is named. Kamonyi’s action is all-crop seed, and fertilizer is named; irrigation does not fire because smallholders are already at 22.4%. Kirehe’s action is irrigation if rainfall passes, and fertilizer is named. Gakenke is seed only. Ngoma is irrigation, with fertilizer named; its seed gap of 12.2 does not fire. Boundaries: Kirehe 14.9 against 15, Muhanga 15.2, Gakenke’s irrigation gap 28.7 against 30, Rwamagana’s seed gap 24.5 against 25, Kirehe’s fertilizer lag 16.9 against 15.",
    )
    pdf_h(pdf, "5. Refusals, and the Season B check")
    pdf_p(
        pdf,
        "Refused: a district store; a maize variety; a satisfaction score; a learned seed-soil-fertilizer score; copying a national rate onto a district; letting Season B or Season C replace Season A; re-estimating the NST2 hectare total; letting rainfall replace a NISR cell. The 2025 plot questionnaire has a production-status reason. It is not a published district table. The Rwanda Soil Information System is a fertilizer recommendation from soil. Maize was still in testing at the November 2024 launch.",
    )
    pdf_p(
        pdf,
        "Season B, released 22 September 2026, from 1,200 segments and 505 large-scale farmers, collected 19 April to 28 June 2026. The publication page: improved seed 20.9% (smallholders 19.4%, large farms 65%); inorganic fertilizer 57.5% (56.7% and 80.2%); organic fertilizer 81.3%; irrigation 13.1%. Workbook maize yield is 1.29 t/ha, not 1.99. Season B irrigation is Table 29, and that sheet’s footer still says SAS 2025. No Season B irrigation cell is shown until that footer is resolved.",
    )
    pdf_h(pdf, "6. Links")
    links = [
        "Hackathon: https://statistics.gov.rw/about/hackathon/2026-hackathon-competition",
        "SAS Season A 2026: https://statistics.gov.rw/sites/default/files/documents/2026-04/SAS%202026%20Season%20A%20_Final%20Report.pdf",
        "SAS Season B page: https://statistics.gov.rw/statistical-publications/agriculture-environment/seasonal-agricultural-survey-season-b-2026",
        "SAS Season B report: https://statistics.gov.rw/sites/default/files/documents/2026-09/SAS%202026B_Main%20report.pdf",
        "SAS Season B summary: https://statistics.gov.rw/sites/default/files/documents/2026-09/SAS%202026B_Summary%20in%20English.pdf",
        "SAS Season B workbook: https://statistics.gov.rw/sites/default/files/documents/2026-09/Tables_2026%20Season%20B.xlsx",
        "SAS 2025 annual: https://statistics.gov.rw/sites/default/files/documents/2025-12/SAS%202025%20Final%20report.pdf",
        "Microdata: https://microdata.statistics.gov.rw/index.php/catalog/124",
        "Plot questionnaire: https://www.microdata.statistics.gov.rw/index.php/catalog/124/download/1114",
        "PSTA 5: https://www.minagri.gov.rw/fileadmin/user_upload/Minagri/Publications/Policies_and_strategies/PSTA_5_Abridged_version_Final.pdf",
        "NST2: https://www.minecofin.gov.rw/fileadmin/user_upload/Minecofin/Publications/STRATEGIES/NST_2/NST2_2024-2029_Abridged.pdf",
        "Umushyikirano resolutions: https://www.umushyikirano.gov.rw/fileadmin/user_upload/Umushyikirano/Imyanzuro/Resolutions_20th_UMUSHYIKIRANO_-_17_Feb_26.pdf",
        "AHS 2024 report: https://statistics.gov.rw/sites/default/files/documents/2025-10/AHS%202024_Main%20report.pdf",
        "AHS 2024 page: https://statistics.gov.rw/data-sources/surveys/Agricultural-Household-Survey/agricultural-household-survey-2024",
        "CHIRPS: https://www.chc.ucsb.edu/data/chirps",
        "RwaSIS: https://www.rab.gov.rw/1-1/news-details/rab-launches-rwanda-digital-soil-information-system-to-improve-agricultural-productivity",
        "SAS Season A 2024, not the base file: https://statistics.gov.rw/sites/default/files/documents/2024-10/SAS_2024%20Season%20A%20Final%20Report%20%282%29_compressed.pdf",
        "GitLab: https://gitlab.com/mukunzijonathan/umusaruro",
        "KT Press, 23 March 2026: https://www.ktpress.rw/2026/03/kagame-criticizes-slow-progress-in-agriculture-urges-shift-from-promises-to-action/",
        "AllAfrica, 29 June 2026, news not Hansard: https://allafrica.com/stories/202606300076.html",
    ]
    for line in links:
        pdf.set_font("Lib", "", 7.5)
        pdf.set_text_color(26, 70, 120)
        pdf.multi_cell(0, 3.6, line, align="L")
        pdf.ln(0.4)
    pdf.ln(1)
    pdf_h(pdf, "7. Phases")
    pdf_table(
        pdf,
        ["Phase", "Done when"],
        [
            ["1. Contract", "Both people accept this document. No screen yet."],
            ["2. Data file", "Second person checks every cell. Nationals reproduce. CHIRPS is computed."],
            ["3. Rules", "Tests: Nyagatare, Kirehe, Kamonyi, Gakenke, Ngoma, one blank."],
            ["4. Django", "Map, district page, method page. Source on every number."],
            ["5. Season B", "Beside Season A. Cannot overwrite it. Table 29 footer resolved."],
            ["6. Submit", "Public repo, live link, disclosure. 30 October 2026, 23:59."],
        ],
        [28, 154],
    )
    pdf.output(PDF)


if __name__ == "__main__":
    build_docx()
    build_pdf()
    print(DOCX)
    print(PDF)
